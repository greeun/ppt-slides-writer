#!/usr/bin/env python3
"""html2pptx.py — deck.html → 네이티브 PPTX 변환기 (python-pptx).

사용법:
    python3 html2pptx.py <deck.html> <out.pptx> [--image-slides]

변환 계약 (references/conversion-rules.md):
  - 캔버스 1280x720px → 슬라이드 12192000x6858000 EMU. 1px = 9525 EMU.
  - 폰트 pt = px x 0.75 (0.5pt 반올림). font-family는 첫 번째 패밀리명.
  - 파싱 대상: section.slide 직계의 .el-text / .el-image / .el-shape / .el-table
    + aside.notes. 이외 요소 발견 시 오류 목록 출력 후 exit 2.
  - .el-text 리스트는 1단만 허용 — li 안의 ul/ol 중첩은 오류(exit 2).
  - .el-text 안의 <a href>는 PPTX run 하이퍼링크로 이관된다 (CTA 클릭 추적 경로).
  - 인라인 style의 var(--토큰)은 <style> :root 값으로 해석한다.
  - font-family 폴백 체인: 블록 인라인 → .el-* 인라인 → <style>의 .el-text/body 규칙.
  - --image-slides: headless Chrome으로 장당 2x PNG 캡처 → full-bleed 이미지 슬라이드
    (100% 비주얼, 텍스트 편집 불가). 이때도 노트는 이관한다.

의존성: python-pptx, beautifulsoup4 (+ --image-slides 시 headless Chrome)
"""

import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile

from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

EMU_PER_PX = 9525
SLIDE_W_EMU = 12192000
SLIDE_H_EMU = 6858000

DEFAULT_FONT_PX = {"h1": 40, "h2": 32, "h3": 24, "p": 18, "li": 18}
ALIGN_MAP = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
             "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
NAMED_COLORS = {"white": "#ffffff", "black": "#000000"}

errors = []
default_family = None  # main()이 <style>의 .el-text/body 규칙에서 추출 — 폰트 폴백 체인 최하단


def emu(px_val):
    return Emu(int(round(px_val * EMU_PER_PX)))


def pt_from_px(px_val):
    return Pt(round(px_val * 0.75 * 2) / 2)


def parse_style(style_str):
    out = {}
    if not style_str:
        return out
    for part in style_str.split(";"):
        if ":" not in part:
            continue
        k, v = part.split(":", 1)
        out[k.strip().lower()] = v.strip()
    return out


def parse_root_vars(style_text):
    out = {}
    m = re.search(r":root\s*\{([^}]*)\}", style_text, re.S)
    if not m:
        return out
    for name, val in re.findall(r"(--[\w-]+)\s*:\s*([^;]+);?", m.group(1)):
        out[name.strip()] = val.strip()
    return out


def resolve_vars(value, root_vars):
    if value is None:
        return value
    for _ in range(3):
        def rep(m):
            name = m.group(1).strip()
            fb = m.group(2)
            return root_vars.get(name, fb.strip() if fb else m.group(0))
        new = re.sub(r"var\(\s*(--[\w-]+)\s*(?:,([^)]*))?\)", rep, value)
        if new == value:
            break
        value = new
    return value


def px(value):
    if value is None:
        return None
    m = re.match(r"^(-?\d+(?:\.\d+)?)px$", str(value).strip())
    return float(m.group(1)) if m else None


def parse_color(value):
    """#hex / rgb() / 이름 → RGBColor 또는 None(transparent 등)."""
    if not value:
        return None
    v = value.strip().lower()
    if v in ("transparent", "none", "inherit"):
        return None
    v = NAMED_COLORS.get(v, v)
    m = re.match(r"^#([0-9a-f]{3})$", v)
    if m:
        v = "#" + "".join(c * 2 for c in m.group(1))
    m = re.match(r"^#([0-9a-f]{6})$", v)
    if m:
        return RGBColor.from_string(m.group(1))
    m = re.match(r"^rgb\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*\)$", v)
    if m:
        return RGBColor(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def first_font_family(value):
    if not value:
        return None
    first = value.split(",")[0].strip().strip("'\"")
    return first or None


def doc_default_family(style_text, root_vars):
    """<style> 규칙에서 문서 기본 font-family 추출 (.el-text 규칙 우선, 다음 body 규칙).

    인라인 미지정 시 브라우저 캐스케이드가 적용하는 폰트를 PPTX run에도 도달시킨다
    — 폴백 최하단 (conversion-rules §2)."""
    found = {}
    for m in re.finditer(r"([^{}]+)\{([^}]*)\}", style_text):
        sel = m.group(1)
        fam = parse_style(m.group(2)).get("font-family")
        if not fam:
            continue
        fam = first_font_family(resolve_vars(fam, root_vars))
        if not fam or fam.startswith("var("):
            continue
        if ".el-text" in sel:
            found["el-text"] = fam
        elif any(tok.strip() == "body" for tok in sel.split(",")):
            found["body"] = fam
    return found.get("el-text") or found.get("body")


def geometry(st, section_idx, kind):
    l, t = px(st.get("left")), px(st.get("top"))
    w, h = px(st.get("width")), px(st.get("height"))
    if None in (l, t, w, h):
        errors.append(f"S{section_idx} {kind}: left/top/width/height(px) 인라인 배치 누락.")
        return None
    return l, t, w, h


def merged_style(el, root_vars, base=None):
    st = dict(base or {})
    own = {k: resolve_vars(v, root_vars) for k, v in parse_style(el.get("style", "")).items()}
    st.update(own)
    return st


def set_bullet(paragraph, on):
    """li → 네이티브 불릿(buChar), 그 외 → buNone."""
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for old in pPr.findall(qn(tag)):
            pPr.remove(old)
    if on:
        pPr.set("marL", "228600")
        pPr.set("indent", "-228600")
        bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
        pPr.append(bu)
    else:
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def add_runs(paragraph, node, ctx):
    """인라인 노드 순회 → run 생성. b/strong=bold, em/i=italic, span(color/font-weight),
    a[href]=하이퍼링크(run.hyperlink — PPTX·PDF 양쪽에서 클릭 가능)."""
    for child in node.children:
        if isinstance(child, Comment):
            continue  # Comment는 NavigableString의 하위 타입 — run으로 새지 않게 차단
        if isinstance(child, NavigableString):
            text = re.sub(r"\s+", " ", str(child))
            if not text.strip():
                continue
            run = paragraph.add_run()
            run.text = text
            _apply_run_style(run, ctx)
        elif isinstance(child, Tag):
            sub = dict(ctx)
            if child.name in ("b", "strong"):
                sub["bold"] = True
            if child.name in ("em", "i"):
                sub["italic"] = True
            if child.name == "br":
                # 줄바꿈은 vertical-tab으로 처리 (python-pptx run 내 개행)
                run = paragraph.add_run()
                run.text = "\v"
                continue
            if child.name == "a":
                href = (child.get("href") or "").strip()
                if href:
                    sub["hyperlink"] = href
            if child.name == "span":
                st = merged_style(child, sub["root_vars"])
                c = parse_color(st.get("color"))
                if c:
                    sub["color"] = c
                if st.get("font-weight") in ("bold", "600", "700", "800", "900"):
                    sub["bold"] = True
                fs = px(st.get("font-size"))
                if fs:
                    sub["size_px"] = fs
            add_runs(paragraph, child, sub)


def _apply_run_style(run, ctx):
    run.font.size = pt_from_px(ctx["size_px"])
    if ctx.get("hyperlink"):
        run.hyperlink.address = ctx["hyperlink"]
    if ctx.get("bold"):
        run.font.bold = True
    if ctx.get("italic"):
        run.font.italic = True
    if ctx.get("color") is not None:
        run.font.color.rgb = ctx["color"]
    if ctx.get("family"):
        run.font.name = ctx["family"]


def block_ctx(block, el_style, root_vars):
    """블록(h1~h3/p/li)의 유효 스타일 컨텍스트 계산 (인라인 → .el-* 인라인 → 기본값)."""
    st = merged_style(block, root_vars, base=el_style)
    size = px(st.get("font-size")) or DEFAULT_FONT_PX.get(block.name, 18)
    color = parse_color(st.get("color"))
    family = first_font_family(st.get("font-family")) or default_family
    bold = st.get("font-weight") in ("bold", "600", "700", "800", "900") \
        or block.name in ("h1", "h2", "h3")
    align = ALIGN_MAP.get(st.get("text-align", "").strip())
    line_height = st.get("line-height")
    return {"size_px": size, "color": color, "family": family, "bold": bold,
            "align": align, "line_height": line_height, "root_vars": root_vars}


def apply_paragraph(paragraph, ctx):
    if ctx.get("align") is not None:
        paragraph.alignment = ctx["align"]
    lh = ctx.get("line_height")
    if lh:
        lh = lh.strip()
        if re.match(r"^\d+(\.\d+)?$", lh):
            paragraph.line_spacing = float(lh)
        elif px(lh):
            paragraph.line_spacing = pt_from_px(px(lh))


def add_text_element(slide, el, idx, root_vars):
    geo = geometry(merged_style(el, root_vars), idx, ".el-text")
    if not geo:
        return
    l, t, w, h = geo
    # 중첩 리스트는 blocks 수집이 중첩 li를 이중 방문해 문단을 오염시킨다 —
    # colspan과 동일한 오류 정책으로 거부 (conversion-rules §4, 부분 변환물 금지)
    for lst in el.find_all(["ul", "ol"]):
        anc = lst.find_parent(["li", "ul", "ol"])
        if anc is not None:
            errors.append(
                f"S{idx} .el-text: 중첩 리스트 미지원 — <{lst.name}>이 <{anc.name}> 내부에 있음. "
                "리스트는 1단만 허용(html-spec §6) — 문단 재구성 또는 장표 분할로 해소할 것.")
            return
    box = slide.shapes.add_textbox(emu(l), emu(t), emu(w), emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    el_style = merged_style(el, root_vars)
    blocks = el.find_all(["h1", "h2", "h3", "p", "li"])
    blocks = [b for b in blocks if b.get_text(strip=True)]
    if not blocks:
        blocks = [el]  # 블록 없이 직접 텍스트만 있는 경우
    first = True
    for block in blocks:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        ctx = block_ctx(block, el_style, root_vars)
        apply_paragraph(para, ctx)
        set_bullet(para, block.name == "li")
        add_runs(para, block, ctx)


def add_image_element(slide, el, idx, root_vars, deck_dir):
    geo = geometry(merged_style(el, root_vars), idx, ".el-image")
    if not geo:
        return
    l, t, w, h = geo
    src = el.get("src", "")
    if re.match(r"^https?://", src, re.I):
        errors.append(f"S{idx} .el-image: 외부 URL 금지 — {src}")
        return
    path = os.path.normpath(os.path.join(deck_dir, src))
    if not os.path.isfile(path):
        errors.append(f"S{idx} .el-image: 파일 없음 — {path}")
        return
    slide.shapes.add_picture(path, emu(l), emu(t), emu(w), emu(h))


def add_shape_element(slide, el, idx, root_vars):
    st = merged_style(el, root_vars)
    geo = geometry(st, idx, ".el-shape")
    if not geo:
        return
    l, t, w, h = geo
    radius = px(st.get("border-radius")) or 0
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius > 0 else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, emu(l), emu(t), emu(w), emu(h))
    if radius > 0 and min(w, h) > 0:
        try:
            shp.adjustments[0] = min(0.5, radius / min(w, h))
        except (IndexError, ValueError):
            pass
    fill_color = parse_color(st.get("background-color") or st.get("background"))
    if fill_color is not None:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill_color
    else:
        shp.fill.background()
    border = st.get("border", "")
    bm = re.match(r"^(\d+(?:\.\d+)?)px\s+solid\s+(.+)$", border)
    if bm and parse_color(bm.group(2)):
        shp.line.color.rgb = parse_color(bm.group(2))
        shp.line.width = pt_from_px(float(bm.group(1)))
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    rot = re.search(r"rotate\(\s*(-?\d+(?:\.\d+)?)deg\s*\)", st.get("transform", ""))
    if rot:
        shp.rotation = float(rot.group(1))
    text = el.get_text(strip=True)
    if text:
        tf = shp.text_frame
        tf.word_wrap = True
        tf.text = text
        fs = px(st.get("font-size")) or 18
        fam = first_font_family(st.get("font-family")) or default_family
        for para in tf.paragraphs:
            for run in para.runs:
                run.font.size = pt_from_px(fs)
                if fam:
                    run.font.name = fam
                c = parse_color(st.get("color"))
                if c:
                    run.font.color.rgb = c


def parse_padding(value):
    """CSS padding 단축 표기 → (top, right, bottom, left) px. 미지정·해석 불가면 None."""
    if not value:
        return None
    parts = [px(v) for v in str(value).split()]
    if not parts or any(v is None for v in parts):
        return None
    if len(parts) == 1:
        t = r = b = l = parts[0]
    elif len(parts) == 2:
        t = b = parts[0]; r = l = parts[1]
    elif len(parts) == 3:
        t, r, b = parts; l = r
    else:
        t, r, b, l = parts[:4]
    return t, r, b, l


def table_css_defaults(style_text, root_vars):
    """<style>의 .el-table 셀 규칙(padding·vertical-align·border)을 추출한다.

    boilerplate는 표 기본 스타일을 스타일시트에 두므로(인라인 아님) 변환기가
    같은 값을 PPTX 셀 여백·테두리로 이관하려면 여기서 읽어야 한다.
    반환: {"cell": {...}, "header": {...}} — 각 값은 CSS 선언 dict."""
    out = {"cell": {}, "header": {}}
    for m in re.finditer(r"([^{}]+)\{([^}]*)\}", style_text):
        sel = m.group(1).strip()
        if ".el-table" not in sel:
            continue
        decl = {k: resolve_vars(v, root_vars) for k, v in parse_style(m.group(2)).items()}
        if not decl:
            continue
        key = "header" if "thead" in sel else ("cell" if re.search(r"\b(th|td)\b", sel) else None)
        if key:
            out[key].update(decl)
    return out


def set_cell_borders(cell, edges):
    """PPTX 셀 테두리 설정 (python-pptx 미지원 — a:lnL/R/T/B 직접 삽입).

    edges: {'L'|'R'|'T'|'B': (color|None, width_px)}. color None 또는 폭 0이면 noFill.
    스키마상 tcPr 자식 순서는 lnL → lnR → lnT → lnB 이므로 역순으로 맨 앞에 넣는다."""
    tcPr = cell._tc.get_or_add_tcPr()
    for edge in ("L", "R", "T", "B"):
        for old in tcPr.findall(qn("a:ln" + edge)):
            tcPr.remove(old)
    for edge in ("B", "T", "R", "L"):
        color, width_px = edges.get(edge, (None, 0))
        ln = tcPr.makeelement(qn("a:ln" + edge), {
            "w": str(max(int(round(width_px * 12700 * 0.75)), 1)),
            "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        if color is None or width_px <= 0:
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        else:
            fill = ln.makeelement(qn("a:solidFill"), {})
            fill.append(fill.makeelement(qn("a:srgbClr"), {"val": str(color)}))
            ln.append(fill)
        tcPr.insert(0, ln)


def border_spec(st, side, base_color):
    """CSS border / border-<side> → (RGBColor|None, width_px). 미지정이면 (None, 0)."""
    raw = st.get("border-" + side) or st.get("border")
    if not raw or raw.strip() in ("none", "0", "0px"):
        return (None, 0)
    width = 1.0
    m = re.search(r"(\d+(?:\.\d+)?)px", raw)
    if m:
        width = float(m.group(1))
    color = parse_color(re.sub(r"(\d+(?:\.\d+)?px|solid|dashed|dotted|none)", "", raw).strip()) or base_color
    return (color, width)


def add_table_element(slide, el, idx, root_vars, css_table=None):
    st = merged_style(el, root_vars)
    geo = geometry(st, idx, ".el-table")
    if not geo:
        return
    l, t, w, h = geo
    rows = el.find_all("tr")
    if not rows:
        errors.append(f"S{idx} .el-table: tr 없음.")
        return
    ncols = max(len(r.find_all(["td", "th"])) for r in rows)
    for r in rows:
        for cell in r.find_all(["td", "th"]):
            if cell.get("colspan") or cell.get("rowspan"):
                errors.append(f"S{idx} .el-table: colspan/rowspan 미지원 — 셀 분해 필요.")
                return
    gframe = slide.shapes.add_table(len(rows), ncols, emu(l), emu(t), emu(w), emu(h))
    table = gframe.table
    # 첫 행 강조·줄무늬 기본값 해제 — 색은 셀 인라인 style이 단일 원천이다
    table.first_row = False
    table.horz_banding = False

    # 열 너비: <colgroup><col style="width:Npx"> (브라우저 fixed layout과 PPTX를 일치시킨다)
    cols = el.find_all("col")
    widths = [px(merged_style(c, root_vars).get("width")) for c in cols][:ncols]
    if len(widths) == ncols and all(v for v in widths):
        scale = w / sum(widths)
        for ci, cw in enumerate(widths):
            table.columns[ci].width = emu(cw * scale)
    # 행 높이: <tr style="height:Npx"> 지정 시 이관 (미지정이면 python-pptx 균등 분배)
    row_px = [px(merged_style(r, root_vars).get("height")) for r in rows]
    if all(v for v in row_px):
        scale = h / sum(row_px)
        for ri, rh in enumerate(row_px):
            table.rows[ri].height = emu(rh * scale)

    css_table = css_table or {"cell": {}, "header": {}}
    table_fs = px(st.get("font-size")) or 18  # 미지정 시 본문 최소 크기(html-spec §7)와 동일
    for ri, r in enumerate(rows):
        cells = r.find_all(["td", "th"])
        row_style = merged_style(r, root_vars)
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            if ci >= len(cells):
                continue
            src_cell = cells[ci]
            is_header = src_cell.name == "th"
            # 스타일 우선순위: <style> 셀 규칙 → (헤더면) thead 규칙 → tr 인라인 → 셀 인라인
            base = dict(css_table.get("cell", {}))
            if is_header:
                base.update(css_table.get("header", {}))
            base.update({k: v for k, v in row_style.items() if k not in ("height",)})
            cst = merged_style(src_cell, root_vars, base=base)
            fam = first_font_family(cst.get("font-family") or st.get("font-family")) or default_family
            cell.text = re.sub(r"\s+", " ", src_cell.get_text(" ", strip=True))

            # 셀 여백 = CSS padding (미지정 시 PPTX 기본값 유지)
            pad = parse_padding(cst.get("padding"))
            if pad:
                ptop, pright, pbottom, pleft = pad
                cell.margin_top = emu(ptop)
                cell.margin_right = emu(pright)
                cell.margin_bottom = emu(pbottom)
                cell.margin_left = emu(pleft)
            # 수직 정렬
            va = (cst.get("vertical-align") or "").strip().lower()
            if va in ("middle", "center"):
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            elif va == "bottom":
                cell.vertical_anchor = MSO_ANCHOR.BOTTOM
            elif va == "top":
                cell.vertical_anchor = MSO_ANCHOR.TOP

            bg = parse_color(cst.get("background-color") or cst.get("background"))
            if bg is not None:
                cell.fill.solid()
                cell.fill.fore_color.rgb = bg
            else:
                cell.fill.background()  # PPTX 테마 기본 채우기 제거 (HTML은 투명)

            # 테두리: CSS 지정분만 그린다. 미지정 변은 noFill — 엑셀식 전면 격자를 만들지 않는다
            line_color = parse_color(cst.get("border-color"))
            set_cell_borders(cell, {
                "T": border_spec(cst, "top", line_color),
                "B": border_spec(cst, "bottom", line_color),
                "L": border_spec(cst, "left", line_color),
                "R": border_spec(cst, "right", line_color),
            })

            for para in cell.text_frame.paragraphs:
                al = ALIGN_MAP.get(cst.get("text-align", "").strip())
                if al is not None:
                    para.alignment = al
                for run in para.runs:
                    run.font.size = pt_from_px(px(cst.get("font-size")) or table_fs)
                    if fam:
                        run.font.name = fam
                    if is_header or cst.get("font-weight") in ("bold", "600", "700", "800", "900"):
                        run.font.bold = True
                    c = parse_color(cst.get("color"))
                    if c:
                        run.font.color.rgb = c


def find_chrome():
    cands = []
    if os.environ.get("CHROME_BIN"):
        cands.append(os.environ["CHROME_BIN"])
    cands.append("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
    cands += ["google-chrome", "chromium"]
    for c in cands:
        if os.path.sep in c:
            if os.path.isfile(c) and os.access(c, os.X_OK):
                return c
        else:
            p = shutil.which(c)
            if p:
                return p
    return None


def capture_slide_pngs(deck_path, sections_count, out_dir):
    """--image-slides: 장당 2x PNG 캡처 (1280x720 → 2560x1440)."""
    chrome = find_chrome()
    if not chrome:
        print("오류: --image-slides에는 headless Chrome이 필요하다. "
              "설치 후 재시도하거나 $CHROME_BIN을 지정할 것.", file=sys.stderr)
        sys.exit(3)
    with open(deck_path, encoding="utf-8") as f:
        src = f.read()
    pngs = []
    for i in range(sections_count):
        # i번째 섹션만 보이도록 나머지를 숨기고, 대상 섹션을 (0,0)에 고정
        hide_css = (
            "<style>section.slide{display:none !important;}"
            f"section.slide:nth-of-type({i + 1})"
            "{display:block !important; position:absolute; left:0; top:0;}</style>"
        )
        injected = src.replace("</head>", hide_css + "</head>", 1) \
            if "</head>" in src else hide_css + src
        tmp_html = os.path.join(out_dir, f"slide_{i + 1:02d}.html")
        png = os.path.join(out_dir, f"slide_{i + 1:02d}.png")
        with open(tmp_html, "w", encoding="utf-8") as f:
            f.write(injected)
        subprocess.run(
            [chrome, "--headless=new", "--disable-gpu", "--no-sandbox",
             "--hide-scrollbars", "--force-device-scale-factor=2",
             "--window-size=1280,720", f"--screenshot={png}",
             "file://" + os.path.abspath(tmp_html)],
            capture_output=True, text=True, timeout=120, check=False,
        )
        if not os.path.isfile(png):
            print(f"오류: S{i + 1} 스크린샷 실패.", file=sys.stderr)
            sys.exit(3)
        pngs.append(png)
    return pngs


def main():
    ap = argparse.ArgumentParser(description="deck.html → 네이티브 PPTX")
    ap.add_argument("deck")
    ap.add_argument("output")
    ap.add_argument("--image-slides", action="store_true",
                    help="장당 PNG 캡처로 full-bleed 이미지 슬라이드 구성 (편집 불가 옵션)")
    args = ap.parse_args()

    deck_path = os.path.abspath(args.deck)
    deck_dir = os.path.dirname(deck_path)
    with open(deck_path, encoding="utf-8") as f:
        soup = BeautifulSoup(f.read(), "html.parser")
    sections = soup.select("section.slide")
    if not sections:
        print("오류: section.slide가 없다.", file=sys.stderr)
        sys.exit(2)
    style_text = "\n".join(s.get_text() for s in soup.find_all("style"))
    root_vars = parse_root_vars(style_text)
    global default_family
    default_family = doc_default_family(style_text, root_vars)
    css_table = table_css_defaults(style_text, root_vars)

    # .slide 기본 배경 (CSS 규칙에서 추출, 섹션 인라인이 우선)
    default_bg = None
    for m in re.finditer(r"([^{}]+)\{([^}]*)\}", style_text):
        if ".slide" in m.group(1):
            body = {k: resolve_vars(v, root_vars)
                    for k, v in parse_style(m.group(2)).items()}
            default_bg = parse_color(body.get("background-color") or body.get("background")) or default_bg

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    blank = None
    for layout in prs.slide_layouts:
        if layout.name.lower() == "blank":
            blank = layout
            break
    if blank is None:
        blank = prs.slide_layouts[6]

    image_paths = None
    tmp_shots = None
    if args.image_slides:
        tmp_shots = tempfile.mkdtemp(prefix=".slide-shots-", dir=deck_dir)
        image_paths = capture_slide_pngs(deck_path, len(sections), tmp_shots)

    for idx, section in enumerate(sections, 1):
        slide = prs.slides.add_slide(blank)
        sec_style = {k: resolve_vars(v, root_vars)
                     for k, v in parse_style(section.get("style", "")).items()}
        bg = parse_color(sec_style.get("background-color") or sec_style.get("background")) or default_bg
        if bg is not None:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = bg

        if args.image_slides:
            slide.shapes.add_picture(image_paths[idx - 1], 0, 0,
                                     Emu(SLIDE_W_EMU), Emu(SLIDE_H_EMU))
        else:
            for child in section.children:
                if not isinstance(child, Tag):
                    continue
                classes = child.get("class") or []
                if child.name == "aside" and "notes" in classes:
                    continue
                if "el-text" in classes:
                    add_text_element(slide, child, idx, root_vars)
                elif "el-image" in classes:
                    add_image_element(slide, child, idx, root_vars, deck_dir)
                elif "el-shape" in classes:
                    add_shape_element(slide, child, idx, root_vars)
                elif "el-table" in classes:
                    add_table_element(slide, child, idx, root_vars, css_table)
                else:
                    errors.append(
                        f"S{idx}: 파싱 집합 외 직계 요소 <{child.name} class={classes}> — "
                        ".el-text/.el-image/.el-shape/.el-table/aside.notes만 허용.")

        # 발표자 노트 이관 (이미지 슬라이드 모드에서도 수행)
        notes = section.find("aside", class_="notes")
        if notes is not None:
            slide.notes_slide.notes_text_frame.text = notes.get_text(" ", strip=True)

    if tmp_shots:
        shutil.rmtree(tmp_shots, ignore_errors=True)

    if errors:
        print("변환 오류 — PPTX를 저장하지 않는다:", file=sys.stderr)
        for e in errors:
            print(f"  - {e}", file=sys.stderr)
        sys.exit(2)

    os.makedirs(os.path.dirname(os.path.abspath(args.output)) or ".", exist_ok=True)
    prs.save(args.output)
    mode = "image-slides" if args.image_slides else "native"
    print(f"PPTX 저장 완료 ({mode}, {len(sections)}장): {args.output}")


if __name__ == "__main__":
    main()
