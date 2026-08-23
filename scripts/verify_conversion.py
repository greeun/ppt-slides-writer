#!/usr/bin/env python3
"""verify_conversion.py — 변환 충실도 검증 (체크 21~26).

사용법:
    python3 verify_conversion.py <deck.html> <deck.pptx> <deck.pdf> \
        --report <verify_report.md>

| # | 체크 | 판정 |
|---|---|---|
| 21 | PPTX 재오픈+장수 | Presentation() 재오픈 성공 AND 슬라이드 수 == 섹션 수 |
| 22 | 텍스트 무손실 | HTML 섹션별 텍스트 노드(공백 정규화) ⊆ 해당 슬라이드 텍스트 프레임 집합 |
| 23 | 노트 이관 | aside.notes 텍스트 == notes_slide 텍스트 (공백 정규화) |
| 24 | 이미지 | 섹션별 img 수 == picture 수, 원본 해상도 ≥ 배치 px(하한), <2x는 WARN |
| 25 | PDF | 페이지 수 == 장수 AND 각 페이지 텍스트 레이어 비어있지 않음 |
| 26 | 스키마 호환 | python-pptx round-trip 재저장 성공 (실제 열람 확인은 사람 게이트 ④) |

주의: --image-slides로 만든 PPTX는 텍스트가 이미지라 22·23(텍스트)·24가 성립하지 않는다.
      검증 대상은 네이티브 PPTX(기본 산출물)다.

의존성: python-pptx, beautifulsoup4, pypdf
"""

import argparse
import os
import re
import struct
import sys
import tempfile

from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

results = []  # (check_no, status, message)  status: PASS | WARN | FAIL


def add(no, status, msg):
    results.append((no, status, msg))


def norm(text):
    return re.sub(r"\s+", " ", text or "").strip()


def image_size(path):
    """PNG/JPEG/GIF 픽셀 크기 파서 (외부 의존성 없이)."""
    with open(path, "rb") as f:
        head = f.read(26)
        if head.startswith(b"\x89PNG\r\n\x1a\n"):
            w, h = struct.unpack(">II", head[16:24])
            return w, h
        if head.startswith(b"GIF8"):
            w, h = struct.unpack("<HH", head[6:10])
            return w, h
        if head.startswith(b"\xff\xd8"):
            f.seek(2)
            while True:
                marker = f.read(2)
                if len(marker) < 2 or marker[0] != 0xFF:
                    return None
                code = marker[1]
                if code in (0xD8, 0xD9):
                    continue
                seg_len = struct.unpack(">H", f.read(2))[0]
                if 0xC0 <= code <= 0xCF and code not in (0xC4, 0xC8, 0xCC):
                    data = f.read(5)
                    h, w = struct.unpack(">HH", data[1:5])
                    return w, h
                f.seek(seg_len - 2, 1)
    return None


def slide_all_text(slide):
    """슬라이드의 모든 텍스트(텍스트 프레임 + 표 셀) 연결."""
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
    return norm(" ".join(parts))


def resolve_vars_factory(style_text):
    root_vars = {}
    m = re.search(r":root\s*\{([^}]*)\}", style_text, re.S)
    if m:
        for name, val in re.findall(r"(--[\w-]+)\s*:\s*([^;]+);?", m.group(1)):
            root_vars[name.strip()] = val.strip()

    def resolve(value):
        if value is None:
            return value
        for _ in range(3):
            new = re.sub(
                r"var\(\s*(--[\w-]+)\s*(?:,([^)]*))?\)",
                lambda mm: root_vars.get(mm.group(1).strip(),
                                         (mm.group(2) or mm.group(0)).strip()),
                value)
            if new == value:
                break
            value = new
        return value
    return resolve


def px(value):
    if value is None:
        return None
    m = re.match(r"^(-?\d+(?:\.\d+)?)px$", str(value).strip())
    return float(m.group(1)) if m else None


def style_dict(el):
    out = {}
    for part in (el.get("style") or "").split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            out[k.strip().lower()] = v.strip()
    return out


def main():
    ap = argparse.ArgumentParser(description="변환 충실도 검증 21~26")
    ap.add_argument("deck_html")
    ap.add_argument("deck_pptx")
    ap.add_argument("deck_pdf")
    ap.add_argument("--report", required=True)
    args = ap.parse_args()

    deck_dir = os.path.dirname(os.path.abspath(args.deck_html))
    with open(args.deck_html, encoding="utf-8") as f:
        soup = BeautifulSoup(f.read(), "html.parser")
    sections = soup.select("section.slide")
    n_sections = len(sections)
    style_text = "\n".join(s.get_text() for s in soup.find_all("style"))
    resolve = resolve_vars_factory(style_text)

    # ---------- 체크 21: PPTX 재오픈 + 장수 ----------
    prs = None
    try:
        prs = Presentation(args.deck_pptx)
        n_slides = len(prs.slides)
        if n_slides == n_sections:
            add(21, "PASS", f"재오픈 성공, 슬라이드 {n_slides} == 섹션 {n_sections}.")
        else:
            add(21, "FAIL", f"장수 불일치 — 슬라이드 {n_slides} vs 섹션 {n_sections}.")
    except Exception as e:  # noqa: BLE001 — 재오픈 실패 자체가 판정 대상
        add(21, "FAIL", f"PPTX 재오픈 실패: {e}")

    # ---------- 체크 22: 텍스트 무손실 ----------
    if prs is not None and len(prs.slides) == n_sections:
        missing = []
        for i, section in enumerate(sections):
            stext = slide_all_text(prs.slides[i])
            for node in section.find_all(string=True):
                parent_names = [p.name for p in node.parents]
                if "aside" in parent_names or "style" in parent_names or "script" in parent_names:
                    continue
                t = norm(str(node))
                if not t:
                    continue
                if t not in stext:
                    missing.append(f"S{i + 1}: \"{t[:60]}\"")
        if missing:
            add(22, "FAIL", f"텍스트 손실 {len(missing)}건: " + " / ".join(missing[:10]))
        else:
            add(22, "PASS", "HTML 텍스트 노드 전부가 슬라이드 텍스트 프레임에 존재 (손실 0).")
    else:
        add(22, "FAIL", "체크 21 실패로 대조 불가.")

    # ---------- 체크 23: 노트 이관 ----------
    if prs is not None and len(prs.slides) == n_sections:
        bad = []
        for i, section in enumerate(sections):
            aside = section.find("aside", class_="notes")
            html_notes = norm(aside.get_text(" ")) if aside else ""
            slide = prs.slides[i]
            pptx_notes = norm(slide.notes_slide.notes_text_frame.text) if slide.has_notes_slide else ""
            if html_notes != pptx_notes:
                bad.append(f"S{i + 1} (HTML {len(html_notes)}자 vs PPTX {len(pptx_notes)}자)")
        if bad:
            add(23, "FAIL", "노트 불일치: " + ", ".join(bad[:10]))
        else:
            add(23, "PASS", "전 장표 노트 일치 (공백 정규화 기준).")
    else:
        add(23, "FAIL", "체크 21 실패로 대조 불가.")

    # ---------- 체크 24: 이미지 ----------
    if prs is not None and len(prs.slides) == n_sections:
        img_fail, img_warn, checked = [], [], 0
        for i, section in enumerate(sections):
            imgs = [el for el in section.find_all("img")
                    if "el-image" in (el.get("class") or [])]
            pics = [sh for sh in prs.slides[i].shapes
                    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
            if len(imgs) != len(pics):
                img_fail.append(f"S{i + 1}: img {len(imgs)} vs picture {len(pics)}")
                continue
            for el in imgs:
                checked += 1
                st = {k: resolve(v) for k, v in style_dict(el).items()}
                w_px, h_px = px(st.get("width")), px(st.get("height"))
                path = os.path.normpath(os.path.join(deck_dir, el.get("src", "")))
                size = image_size(path) if os.path.isfile(path) else None
                if size is None:
                    img_warn.append(f"S{i + 1} {el.get('src')}: 해상도 판독 불가")
                    continue
                nw, nh = size
                if w_px and h_px:
                    if nw < w_px or nh < h_px:
                        img_fail.append(
                            f"S{i + 1} {el.get('src')}: 원본 {nw}x{nh} < 배치 {w_px:.0f}x{h_px:.0f}")
                    elif nw < 2 * w_px or nh < 2 * h_px:
                        img_warn.append(
                            f"S{i + 1} {el.get('src')}: 원본 {nw}x{nh} < 배치 2x — 인쇄 선명도 저하 가능")
        if img_fail:
            add(24, "FAIL", "; ".join(img_fail[:10]))
        elif img_warn:
            add(24, "WARN", f"수량 일치, 해상도 경고 {len(img_warn)}건: " + "; ".join(img_warn[:5]))
        else:
            add(24, "PASS", f"이미지 수량 일치, 해상도 하한 충족 (검사 {checked}건).")
    else:
        add(24, "FAIL", "체크 21 실패로 대조 불가.")

    # ---------- 체크 25: PDF ----------
    try:
        reader = PdfReader(args.deck_pdf)
        n_pages = len(reader.pages)
        empty_pages = [i + 1 for i, p in enumerate(reader.pages)
                       if not (p.extract_text() or "").strip()]
        if n_pages != n_sections:
            add(25, "FAIL", f"PDF 페이지 {n_pages} != 장수 {n_sections}.")
        elif empty_pages:
            add(25, "FAIL", f"텍스트 레이어 빈 페이지: {empty_pages} — 검색 불가 인쇄본.")
        else:
            add(25, "PASS", f"페이지 {n_pages} == 장수, 전 페이지 텍스트 레이어 존재.")
    except Exception as e:  # noqa: BLE001 — PDF 판독 실패 자체가 판정 대상
        add(25, "FAIL", f"PDF 판독 실패: {e}")

    # ---------- 체크 26: 스키마 호환 (round-trip) ----------
    if prs is not None:
        try:
            fd, tmp = tempfile.mkstemp(suffix=".pptx", dir=deck_dir)
            os.close(fd)
            prs.save(tmp)
            Presentation(tmp)
            os.remove(tmp)
            add(26, "PASS", "round-trip 재저장·재오픈 성공 — 스키마 유효 근사. "
                            "파워포인트/Keynote 실제 열람 확인은 최종 게이트(사람 ④) 항목.")
        except Exception as e:  # noqa: BLE001 — round-trip 실패 자체가 판정 대상
            add(26, "FAIL", f"round-trip 실패: {e}")
    else:
        add(26, "FAIL", "체크 21 실패로 round-trip 불가.")

    # ---------- 보고서 ----------
    n_fail = len([r for r in results if r[1] == "FAIL"])
    n_warn = len([r for r in results if r[1] == "WARN"])
    lines = [
        "# Verify Report — 변환 충실도 21~26", "",
        f"- HTML: `{args.deck_html}` ({n_sections}장)",
        f"- PPTX: `{args.deck_pptx}` / PDF: `{args.deck_pdf}`",
        f"- 결과: **FAIL {n_fail} / WARN {n_warn} / PASS {len(results) - n_fail - n_warn}**",
        "",
        "| # | 판정 | 내용 |",
        "|---|---|---|",
    ]
    for no, status, msg in sorted(results, key=lambda x: x[0]):
        lines.append(f"| {no} | {status} | {msg.replace('|', '\\|')} |")
    lines += ["", "> FAIL 존재 시 P0로 수정 후 재변환·재검증한다. 전체 통과 전 최종 게이트 진입 금지."]
    os.makedirs(os.path.dirname(os.path.abspath(args.report)) or ".", exist_ok=True)
    with open(args.report, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")

    print(f"verify: FAIL {n_fail} / WARN {n_warn} → {args.report}")
    sys.exit(1 if n_fail else 0)


if __name__ == "__main__":
    main()
